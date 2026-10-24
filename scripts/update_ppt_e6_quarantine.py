"""Withdraw invalid official E6 claims in the existing results deck."""
import argparse
import hashlib
import os
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def set_text(shape, expected, new):
    assert expected in shape.text, f"missing expected text: {expected[:60]}"
    frame = shape.text_frame
    first_paragraph = frame.paragraphs[0]
    runs = first_paragraph.runs
    assert runs, shape.name
    runs[0].text = new
    for run in list(first_paragraph.runs)[1:]:
        first_paragraph._p.remove(run._r)
    for paragraph in list(frame.paragraphs)[1:]:
        frame._txBody.remove(paragraph._p)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--ppt", required=True)
    parser.add_argument("--image", required=True)
    parser.add_argument("--backup", required=True)
    args = parser.parse_args()
    ppt = Path(args.ppt)
    image = Path(args.image)
    backup = Path(args.backup)
    backup.parent.mkdir(parents=True, exist_ok=True)
    if not backup.exists():
        shutil.copy2(ppt, backup)
    prs = Presentation(str(ppt))
    assert len(prs.slides) == 19

    found_e6_progress = False
    for shape in prs.slides[1].shapes:
        if not shape.has_table:
            continue
        for row in shape.table.rows:
            if row.cells[0].text.strip() == "E6 灾难遗忘（C6b）":
                set_text(row.cells[1], "✅ 87 格全谱",
                         "controlled 30 格；official 重算中")
                set_text(row.cells[2], "51 ncRNA-起",
                         "controlled causal NLL 30 格保留；official MLM 旧值撤回，_mlm15fix 30 格重跑中")
                found_e6_progress = True
    assert found_e6_progress, "E6 progress row not found"

    s17 = prs.slides[16]
    set_text(s17.shapes[2],
             "官方谱（参数量升序）：LoRA 遗忘恒小于全参且随规模衰减",
             "controlled：30M full +26.2 / LoRA +3.4（尾部约 7.8×）；official 旧值撤回、重算中")
    set_text(s17.shapes[3],
             "RiNALMo 官方三档",
             "证据边界：本页仅含 controlled causal NLL；official MLM 旧值因评估实现错误隔离 · _mlm15fix 30 格重跑中 · 正=遗忘 负=反向增益")
    pictures = [s for s in s17.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    pic = pictures[0]
    geometry = (pic.left, pic.top, pic.width, pic.height)
    pic._element.getparent().remove(pic._element)
    s17.shapes.add_picture(str(image), *geometry)

    s18 = prs.slides[17]
    set_text(s18.shapes[9], "87 格遗忘矩阵",
             s18.shapes[9].text.replace("87 格遗忘矩阵",
                                        "controlled E6 30 格（official 30 格重算中）"))

    s19 = prs.slides[18]
    set_text(s19.shapes[28], "灾难遗忘（E6）",
             "资源（E5）：LoRA≈frozen 显存、仅 650M full 超 24GB；遗忘（E6，controlled）：30M 危险带、LoRA 尾部约 7.8×——决策树：小数据域→full+调LR；需保骨干/预算敏感→LoRA（official MLM 重算中）")
    set_text(s19.shapes[30], "LoRA 在遗忘/资源敏感场景全胜",
             "论文主线一句话：微调 RNA LM 时，先修学习率，再按「数据量 × 切分方式 × 参数预算」选策略——full-FT 在小数据域全胜；controlled 遗忘/资源敏感场景优先 LoRA（official MLM 重算中）")

    tmp = ppt.with_name(ppt.stem + ".tmp.pptx")
    prs.save(str(tmp))
    with zipfile.ZipFile(tmp) as archive:
        assert archive.testzip() is None
        media_hashes = {hashlib.md5(archive.read(name)).hexdigest()
                        for name in archive.namelist() if name.startswith("ppt/media/")}
    image_hash = hashlib.md5(image.read_bytes()).hexdigest()
    assert image_hash in media_hashes, "new controlled figure is not embedded"
    check = Presentation(str(tmp))
    assert len(check.slides) == 19
    all_text = "\n".join(shape.text for slide in check.slides for shape in slide.shapes
                         if hasattr(shape, "text"))
    assert "0.118→0.026" not in all_text
    assert "87 格遗忘矩阵" not in all_text
    assert "controlled E6 30 格" in all_text
    assert "official MLM 重算中" in all_text
    table_text = "\n".join(
        cell.text for slide in check.slides for shape in slide.shapes
        if shape.has_table for row in shape.table.rows for cell in row.cells)
    assert "87 格全谱" not in table_text
    assert "0.157→0.026" not in table_text
    assert "controlled causal NLL 30 格保留" in table_text
    os.replace(tmp, ppt)
    print("PPT_E6_QUARANTINE_OK", {"slides": [2, 17, 18, 19], "backup": str(backup),
                                  "embedded_md5": image_hash})


if __name__ == "__main__":
    main()
