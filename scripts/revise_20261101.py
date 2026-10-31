from hashlib import sha256
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

WORK = Path(__file__).resolve().parent
ROOT = WORK.parent.parent
SRC = ROOT / "ppt/RNA_LM微调策略评测-导师讨论修订版_20261030_八任务扩展.pptx"
OUT = ROOT / "ppt/RNA_LM微调策略评测-导师讨论修订版_20261101_导师反馈修订.pptx"
SHUFFLE_FIG = ROOT / "交接文档/绘图数据/current/figure_source_data_20261028/figures/ncrna_shuffle_audit.png"
EXPECTED_SHA = "f4277d59af9f1c2436c3e8e0b99feffffaf0af13443b18088fb3b7984c1d7025"


def shape_by_name(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise KeyError(name)


def set_text(slide, name, text):
    frame = shape_by_name(slide, name).text_frame
    paragraph = frame.paragraphs[0]
    assert paragraph.runs, name
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)
    for extra in list(frame.paragraphs[1:]):
        extra._p.getparent().remove(extra._p)
    paragraph.line_spacing = 1.5


def replace_picture(slide, name, path):
    old = shape_by_name(slide, name)
    left, top, width = old.left, old.top, old.width
    old._element.getparent().remove(old._element)
    with Image.open(path) as image:
        w, h = image.size
    picture = slide.shapes.add_picture(str(path), left, top, width=width)
    picture.height = int(width * h / w)
    picture.name = name


def remove_answered_annotations(prs):
    removed = []
    for index, slide in enumerate(prs.slides, 1):
        for shape in list(slide.shapes):
            text = getattr(shape, "text", "").strip()
            if shape.left < 0 and text:
                removed.append((index, shape.name, text))
                shape._element.getparent().remove(shape._element)
    return removed


def main():
    assert sha256(SRC.read_bytes()).hexdigest() == EXPECTED_SHA
    prs = Presentation(str(SRC))
    assert len(prs.slides) == 30

    set_text(prs.slides[1], "PurposeNote",
             "八任务的目的：检验结论是否随任务粒度、分类/回归、类别不平衡和 split 改变，区分普遍规律与任务边界；不是堆实验数量。最终输出“何时用何种策略”的条件化配方。")

    slide = prs.slides[4]
    set_text(slide, "TextBox 2", "结论 ①  ncRNA 旧“家族崩溃”可由样本顺序伪影解释")
    set_text(slide, "TextBox 3",
             "只改变每 epoch 是否重排训练样本：shuffle ACC=0.895，no-shuffle=0.096；旧 family 结果无法区分分布偏移与优化伪影，故隔离。")
    set_text(slide, "TextBox 5",
             "读图：shuffle=每轮重排训练样本，不改标签或 split；no-shuffle=保留同类成块顺序。中图红线仅 epoch0–3，是 best0+patience3 的完整停止曲线；结论①用固定 LR 对照，逐任务 validation 选 LR 见结论②。")
    replace_picture(slide, "Picture 6", SHUFFLE_FIG)

    slide = prs.slides[5]
    set_text(slide, "TextBox 2", "结论 ①a-2  modification：family 仍高 AUC，但非“免疫”")
    set_text(slide, "TextBox 3",
             "这是序列级 12 标签任务。已测非 650M 模型的 random−family 差值为 −0.211~+0.058，方向与幅度随模型/策略变化，不能称对切分免疫。")
    set_text(slide, "TextBox 5",
             "读图：六柱为 frozen/LoRA/full × random/family 的三种子均值，重点看同模型同策略的两柱差。LoRA 多数下降约 0.04–0.06；mega-full 的 family 反而高 0.211。650M 未闭合为 NA。")

    slide = prs.slides[6]
    set_text(slide, "TextBox 2", "结论 ①a-3  MRL：family 普遍下降，RiNALMo 三档性能接近")
    set_text(slide, "TextBox 3",
             "micro/mega/650M 的 LoRA random=0.804/0.804/0.793，family=0.711/0.725/0.710；规模差异小，符合当前 MRL 任务的规模收益饱和，并非异常。")
    set_text(slide, "TextBox 5",
             "读图：三档同属 RiNALMo 家族，接近结果只说明当前数据、指标与协议下规模增益有限，不能推广到所有任务。650M frozen random 为 n=2 斜线，full 等未闭合项为 NA；family 降幅按模型与策略分别比较。")

    composite = ROOT / "交接文档/绘图数据/current/figure_source_data_20261101_classification_composite/figures"
    slide = prs.slides[8]
    set_text(slide, "TextBox 2", "分类机制：ncRNA 表示与 full-test 性能同图对照")
    set_text(slide, "TextBox 3",
             "左：pre/LoRA/full 的 CLS 联合 t-SNE；右：同 seed17 checkpoint 的 Frozen/LoRA/Full ACC。聚类增强与泛化提升相关但不等同。")
    set_text(slide, "TextBox 5",
             "读图：上 random、下 family；性能来自 full-test，投影只用于显示。corrected seed17 下 LoRA/full 均保持高 ACC，旧“分类头坍缩”不作因果结论。其余分类任务沿用同版式，回归机制页保留。")
    picture = shape_by_name(slide, "图片 7")
    picture.left, picture.top, picture.width = Inches(0.95), Inches(1.42), Inches(11.4)
    replace_picture(slide, "图片 7", composite / "ncrna_classification_composite.png")

    slide = prs.slides[10]
    set_text(slide, "TextBox 2", "分类机制：modification 的多标签表示与性能")
    set_text(slide, "TextBox 3",
             "这是序列级 12 标签分类，m6A 为索引 7。左看表示，右看同 seed17 的 full-test m6A AUC；family 高 AUC 不等于对 split 免疫。")
    set_text(slide, "TextBox 5",
             "读图：上 random、下 family；红色为 m6A 阳性，灰色为阴性。性能与投影使用同一 checkpoint，但不以聚类形态推断因果；多标签不转成 softmax 单分类。")
    picture = shape_by_name(slide, "图片 7")
    picture.left, picture.top, picture.width = Inches(0.95), Inches(1.42), Inches(11.4)
    replace_picture(slide, "图片 7", composite / "modification_classification_composite.png")

    slide = prs.slides[11]
    set_text(slide, "TextBox 2", "分类机制：CRISPR 的表示与 PR-AUC 对照")
    set_text(slide, "TextBox 3",
             "极不平衡分类任务。左为 pre/LoRA/full 表示，右为同 seed17 的 full-test PR-AUC；当前 full 低于 LoRA，需结合 LR 与优化分析。")
    set_text(slide, "TextBox 5",
             "读图：上 random、下 family；红色为阳性。正例富集只用于投影，性能始终在完整 test 上计算；family 是 pair-sequence cluster，不宣称 guide-disjoint。")
    picture = shape_by_name(slide, "图片 7")
    picture.left, picture.top, picture.width = Inches(0.9), Inches(1.42), Inches(11.5)
    replace_picture(slide, "图片 7", composite / "crispr_classification_composite.png")

    slide = prs.slides[12]
    set_text(slide, "TextBox 2", "分类机制：SSP 的碱基对表示与 pair-F1")
    set_text(slide, "TextBox 3",
             "pair-level 分类。左为 PairHead 输入表示，右为同 seed17 的 full-test pair-F1；random 上 LoRA 略高，family 上 full 更高，仅限该种子。")
    set_text(slide, "TextBox 5",
             "读图：上 random、下 family；红色为 paired，灰色为 unpaired。正负各半只用于投影，性能使用完整 pair 集合与各状态 validation threshold，不用 sampled F1。")
    picture = shape_by_name(slide, "图片 7")
    picture.left, picture.top, picture.width = Inches(0.95), Inches(1.42), Inches(11.4)
    replace_picture(slide, "图片 7", composite / "ssp_classification_composite.png")

    slide = prs.slides[17]
    set_text(slide, "TextBox 5",
             "读图：3e-4 是 runner 的 argparse 工程缺省，不是搜参最优；tuned 由 seed101 在 validation 选档，正式种子为 17/29/43。MRL −0.112 已经独立 CUDA 重评一致：预测近常数（SD=6.25×10⁻⁷），属于 micro-full@3e-4 输出坍缩，不是评测算错；tuning 格不进正式均值。")

    slide = prs.slides[19]
    set_text(slide, "TextBox 2", "结论 ④  数据量等价：small-full(n) vs large-LoRA(m)")
    set_text(slide, "TextBox 3",
             "同家族 33M-full@1e-5 vs 148M-LoRA@3e-4；虚线=|Δ|≤0.02 且 95% CI 覆盖 0。4/8 任务可判，其余标待补或低信号。")
    set_text(slide, "TextBox 5",
             "读图：ncRNA 与 MRL 有等价候选；SSP 当前未观测到，modification 仅低预算有候选。与⑤共用曲线：⑤看模型内数据量效应，④看跨规模预算匹配；不插值、不跨家族。")
    picture = shape_by_name(slide, "Picture 6")
    picture.left, picture.top, picture.width = Inches(1.75), Inches(1.52), Inches(9.8)
    replace_picture(slide, "Picture 6",
                    ROOT / "交接文档/绘图数据/current/figure_source_data_20261030_eight_task_v14/scripts/equivalence_8task.png")

    slide = prs.slides[21]
    set_text(slide, "TextBox 2", "结论 ⑤  数据量 scaling：标注量如何改变策略收益")
    set_text(slide, "TextBox 3",
             "同模型内随数据量比较策略；当前展示 33M-full/33M-LoRA/148M-LoRA，Frozen 七档已注册待补。4/8 任务有完整曲线，其他不补零。")
    set_text(slide, "TextBox 5",
             "读图：与④同一批数据、两种问法——本页看曲线形状、平台与翻转；④再找 small-full(n)≈large-LoRA(m)。CRISPR 低信号，SpliceAI/降解/SSI 待补。")

    slide = prs.slides[22]
    set_text(slide, "TextBox 2", "结论⑤补充  大模型覆盖与证据边界")
    set_text(slide, "TextBox 4",
             "数据：最新 best-val 台账 + MRL/modification 聚合 + B20 阶段表（仅补 frozen）；650M 大格在自然空出的独占显存窗口补测，已有 LoRA/frozen 保留，full 等未闭合项为 NA，不补零。")
    set_text(slide, "ReadNote",
             "读图：格内数字为已完成种子数，0 表示未闭合而非 0 分。148M 在 ncRNA/CRISPR/modification/MRL/SSP 六臂三种子齐，其余任务继续补；650M 只展示已有格，不提前写规模律。")

    slide = prs.slides[23]
    set_text(slide, "TextBox 2", "结论④补充  MRL 高预算：legacy 局部对照")
    set_text(slide, "TextBox 5",
             "读图：full@3000=0.621、LoRA@10000=0.687，并非等价点。本页是 legacy 历史对照，只用于说明④的预算匹配问题，不与当前 best-val 图混算。")

    slide = prs.slides[24]
    set_text(slide, "TextBox 2", "结论④补充  MRL 低预算：legacy 数据收益对照")
    set_text(slide, "TextBox 5",
             "读图：full@1000=0.345、LoRA@3000=0.646，差距明显，不能称等价。保留 frozen 轨迹作历史参照；当前等价候选统一由同协议逐 seed 配对与 CI 判定。")

    slide = prs.slides[25]
    set_text(slide, "TextBox 2", "结论 ⑥  成本：时间×性能的逐任务比较")

    slide = prs.slides[26]
    set_text(slide, "TextBox 2", "结论⑥补充  成本：显存峰值与规模边界")

    slide = prs.slides[27]
    set_text(slide, "TextBox 2", "结论 ⑦  遗忘：controlled RNA-Sc 的 30M 峰值")
    set_text(slide, "TextBox 3",
             "A 图纵轴=ΔNLL(post−pre)，正=遗忘、负=反向增益；点=逐 seed，线=均值。30M 最差种子 full=26.2、LoRA=3.4，尾部比约 7.8×。")
    set_text(slide, "TextBox 5",
             "证据边界：最新 controlled causal held-out NLL 共 30 格、每组 3 seeds；official MLM 使用独立 fixed-mask 协议，不与本图混合。650M LoRA 余格完成后再更新其覆盖，不影响本 controlled 图。")
    replace_picture(slide, "Picture 6",
                    ROOT / "交接文档/绘图数据/current/figure_source_data_20261101_e6/e6_controlled_valid_20261101.png")

    set_text(prs.slides[28], "TextBox 10",
             "证据范围：八任务用于检验结论的任务边界，不是堆数量；SpliceAI random 三种子机制已闭合，family 待最后源格；④/⑤共用同源曲线，当前仅部分面板可判；时间/显存为⑥、遗忘为⑦，未闭合项不作最终结论。")
    set_text(prs.slides[29], "TextBox 21", "对应结论④/⑤")
    set_text(prs.slides[29], "TextBox 26", "对应结论①/机制")

    slide_ids = prs.slides._sldIdLst
    obsolete = slide_ids[20]
    prs.part.drop_rel(obsolete.rId)
    slide_ids.remove(obsolete)
    moving = [slide_ids[22], slide_ids[23]]
    for element in moving:
        slide_ids.remove(element)
    slide_ids.insert(20, moving[0])
    slide_ids.insert(21, moving[1])
    assert len(prs.slides) == 29

    removed = remove_answered_annotations(prs)
    prs.save(str(OUT))
    print("saved", OUT)
    print("slides", len(prs.slides), "annotations_removed", len(removed))
    for item in removed:
        print("removed", item[0], item[1])


if __name__ == "__main__":
    main()
